"""Padrão visual dos documentos do ArcSmith: Brandbook Cloud Target - Modern Work & IA.

Uso no Code Interpreter:
    import sys; sys.path.append('<pasta scripts desta skill>')
    from brand_docx import *
    d = new_document('Título', 'Subtítulo | Cliente', 'Cliente', 'dd/mm/aaaa', 'Consultor')
    key_points(d, ['...']); h2(d, '01', 'Seção'); body(d, 'Texto'); field(d, 'Rótulo', 'Texto')
    table(d, ['Col'], [['valor']], color_cols=(0,)); bullets(d, ['...']); copy_block(d, 'texto literal')
    d.save('Documento.docx')
Requer python-docx. Os logos ficam em ../assets.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

ASSETS = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'assets')
LOGO_DARK = os.path.join(ASSETS, 'logo_ct_colorido.png')   # fundo escuro (capa)
LOGO_LIGHT = os.path.join(ASSETS, 'logo_ct_preto.png')     # fundo claro (cabeçalho)

# Brandbook Cloud Target - Modern Work & IA
DEEP, PURPLE, NEON, BLACK = '3A0E4D', '7E1C9F', 'C53DFF', '06070B'
N1, N2, N3, WHITE = '363636', '808080', 'D3D3D3', 'FFFFFF'
F_HERO, F_SUB, F_BODY, F_LIGHT = 'Poppins ExtraBold', 'Poppins SemiBold', 'Poppins', 'Poppins Light'
CPLX = {'Baixa': NEON, 'Média': PURPLE, 'Alta': DEEP, 'Viável em projeto': N2}
TOWER = 'MODERN WORK & IA'
LABELS = {'pt': {'key': 'PONTOS-CHAVE', 'by': 'Elaborado por', 'page': 'Página '},
          'en': {'key': 'KEY POINTS', 'by': 'Prepared by', 'page': 'Page '}}


def rgb(h): return RGBColor.from_string(h)


def font(run, name=F_BODY, size=10.5, color=N1, bold=False):
    run.font.name = name; run.font.size = Pt(size); run.font.color.rgb = rgb(color); run.bold = bold
    rpr = run._element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rf.set(qn(a), name)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor); tcPr.append(shd)


def borders(cell, left=None, others=None, size='24'):
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{side}')
        col = left if side == 'left' else others
        if col: e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), size if side == 'left' else '4'); e.set(qn('w:color'), col)
        else: e.set(qn('w:val'), 'nil')
        b.append(e)
    tcPr.append(b)


def table_borders(t, color=N3):
    tblPr = t._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), color); b.append(e)
    tblPr.append(b)


def cell_margins(cell, cm=0.35):
    tcPr = cell._tc.get_or_add_tcPr(); m = OxmlElement('w:tcMar')
    for side in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:w'), str(int(cm * 567))); e.set(qn('w:type'), 'dxa'); m.append(e)
    tcPr.append(m)


def para(container, text='', name=F_BODY, size=10.5, color=N1, bold=False, align=None, before=0, after=4, spacing=1.15):
    p = container.add_paragraph() if hasattr(container, 'add_paragraph') else container
    pf = p.paragraph_format; pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = spacing
    if align is not None: p.alignment = align
    if text: font(p.add_run(text), name, size, color, bold)
    return p


def new_document(title, subtitle, client, date, consultant, lang='pt', tower=TOWER):
    L = LABELS[lang]
    d = Document()
    sec = d.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin = sec.bottom_margin = Cm(2.2); sec.left_margin = sec.right_margin = Cm(2.2)
    sec.different_first_page_header_footer = True
    st = d.styles['Normal']; st.font.name = F_BODY; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), F_BODY)
    # header (pages 2+) e rodapés: tabelas de 2 colunas sem borda
    def two_col(container, left_fn, right_fn):
        tb = container.add_table(rows=1, cols=2, width=Cm(16.6))
        tb.columns[0].width = Cm(8.3); tb.columns[1].width = Cm(8.3)
        l, r = tb.cell(0, 0), tb.cell(0, 1)
        for c in (l, r): borders(c); c.paragraphs[0].paragraph_format.space_after = Pt(0)
        left_fn(l.paragraphs[0]); r.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT; right_fn(r.paragraphs[0])
        return tb
    hdr = sec.header; hdr.paragraphs[0].text = ''
    two_col(hdr, lambda p: p.add_run().add_picture(LOGO_LIGHT, width=Cm(3.2)),
            lambda p: font(p.add_run('CLOUD TARGET  ·  MODERN WORK & IA'), F_LIGHT, 7.5, N2))
    def page_num(p):
        font(p.add_run(L['page']), F_LIGHT, 7.5, N2)
        fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
        rr = OxmlElement('w:r'); rpr = OxmlElement('w:rPr')
        rf = OxmlElement('w:rFonts'); rf.set(qn('w:ascii'), F_LIGHT); rf.set(qn('w:hAnsi'), F_LIGHT); rpr.append(rf)
        sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '15'); rpr.append(sz)
        col = OxmlElement('w:color'); col.set(qn('w:val'), N2); rpr.append(col)
        rr.append(rpr); tt = OxmlElement('w:t'); tt.text = '1'; rr.append(tt); fld.append(rr); p._p.append(fld)
    for ft in (sec.footer, sec.first_page_footer):
        two_col(ft, lambda p: font(p.add_run(f'Cloud Target  |  {client}'), F_LIGHT, 7.5, N2), page_num)
    # cover banner (fundo escuro + logo colorido + acento neon)
    t = d.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0); shade(c, BLACK); borders(c); cell_margins(c, 0.9)
    p = c.paragraphs[0]; p.add_run().add_picture(LOGO_DARK, width=Cm(5.2))
    para(c, tower, F_SUB, 8.5, NEON, before=18, after=2)
    para(c, title, F_HERO, 26, WHITE, before=0, after=2, spacing=1.0)
    lp = para(c, '', before=0, after=6)
    pPr = lp._p.get_or_add_pPr(); bd = OxmlElement('w:pBdr'); bt = OxmlElement('w:bottom')
    bt.set(qn('w:val'), 'single'); bt.set(qn('w:sz'), '18'); bt.set(qn('w:color'), PURPLE); bd.append(bt); pPr.append(bd)
    para(c, subtitle, F_SUB, 12, N3, before=6, after=2)
    para(c, f"{date}   |   {L['by']}: {consultant}  ·  Cloud Target", F_LIGHT, 8.5, N3, before=10, after=4)
    para(d, '', after=8)
    d._ct_lang = lang
    return d


def key_points(d, items):
    t = d.add_table(rows=1, cols=1); c = t.cell(0, 0); borders(c, left=NEON, others=N3, size='36'); cell_margins(c, 0.4)
    p = c.paragraphs[0]; font(p.add_run(LABELS[getattr(d, '_ct_lang', 'pt')]['key']), F_SUB, 9, PURPLE)
    for it in items:
        q = c.add_paragraph(style='List Bullet'); q.paragraph_format.space_after = Pt(2)
        font(q.add_run(it), F_BODY, 10, N1)
    para(d, '', after=6)


def h2(d, num, text):
    p = para(d, '', before=14, after=4)
    font(p.add_run(f'{num}  '), F_HERO, 14, PURPLE); font(p.add_run(text), F_HERO, 14, DEEP)


def h3(d, text): para(d, text, F_SUB, 11.5, DEEP, before=8, after=2)


def body(d, text): para(d, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY)


def field(d, label, text):
    p = para(d, '', align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    font(p.add_run(label + ': '), F_SUB, 10.5, DEEP); font(p.add_run(text))


def bullets(d, items, style='List Bullet'):
    for s in items:
        p = d.add_paragraph(style=style); p.paragraph_format.space_after = Pt(2); font(p.add_run(s))


def table(d, header, rows, color_cols=()):
    t = d.add_table(rows=1, cols=len(header)); table_borders(t)
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; shade(c, DEEP); cell_margins(c, 0.15)
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
        font(c.paragraphs[0].add_run(h), F_SUB, 9.5, WHITE)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cell_margins(cells[i], 0.15)
            cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
            col = CPLX.get(v, N1) if i in color_cols else N1
            font(cells[i].paragraphs[0].add_run(v), F_SUB if col != N1 else F_BODY, 9.5, col)
    para(d, '', after=4)


def copy_block(d, text):
    t = d.add_table(rows=1, cols=1); c = t.cell(0, 0); shade(c, BLACK); borders(c, left=NEON, size='36'); cell_margins(c, 0.35)
    lines = text.split('\n')
    first = True
    for ln in lines:
        p = c.paragraphs[0] if first else c.add_paragraph(); first = False
        p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
        font(p.add_run(ln if ln else ' '), F_BODY, 9, WHITE)
    para(d, '', after=4)


